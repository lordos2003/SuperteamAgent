"""Тесты Excel-отчёта (п.22 задания).

Проверяется главное правило: на листе ``AVAILABLE`` не может быть ни одной
ссылки на задание, которое не прошло все фильтры существующего pipeline.
"""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from superteam_agent import report
from superteam_agent.config import (
    UNKNOWN,
    VERIFIED_CLOSED,
    VERIFIED_EXPIRED,
    VERIFIED_HUMAN_ONLY,
    VERIFIED_OPEN,
    VERIFIED_WINNER_ANNOUNCED,
)
from superteam_agent.excel_report import (
    AVAILABLE_COLUMNS,
    EXCLUDED_COLUMNS,
    SHEET_AVAILABLE,
    SHEET_EXCLUDED,
    SHEET_SUMMARY,
    SHEET_UNKNOWN,
    UNKNOWN_COLUMNS,
    days_left,
    sort_excluded,
    write_excel_report,
)
from superteam_agent.runner import is_verified_open_listing

CARD_URL = "https://superteam.fun/earn/listing/slug-1"
FUTURE = "2030-01-01T00:00:00.000Z"
PAST = "2020-01-01T00:00:00.000Z"
NOW = datetime(2026, 9, 15, 9, 0, tzinfo=timezone.utc)


def make_entry(**overrides) -> dict:
    """Запись пайплайна (как в superteam_results.json)."""
    entry = {
        "slug": "slug-1",
        "title": "Fix the CSV export bug",
        "card_url": CARD_URL,
        "url": CARD_URL,
        "reward": "500 USDC",
        "reward_amount": 500,
        "reward_currency": "USDC",
        "token": "USDC",
        "payment_type": "CRYPTO",
        "deadline": FUTURE,
        "deadline_utc": "01 Jan 2030 00:00 UTC",
        "deadline_confirmed": True,
        "deadline_passed": False,
        "has_winners": False,
        "agent_access": "AGENT_ALLOWED",
        "agent_access_unknown": False,
        "region": "Global",
        "eligibility_status": "ELIGIBLE",
        "priority": "HIGH",
        "score": 82,
        "verification_status": VERIFIED_OPEN,
        "financial_risk": "LOW",
        "financial_risk_state": "confirmed_safe",
        "final_decision": "CANDIDATE",
        "exclusion_reasons": [],
        "description": "Build a small Python utility that exports the report data to CSV.",
        "type": "bounty",
        "submissions": 12,
        "difficulty": "MEDIUM",
    }
    entry.update(overrides)
    return entry


def build_workbook(tmp_path: Path, entries: list[dict], available: list[dict]) -> Path:
    """Собрать Excel-отчёт так, как это делает runner (через существующий pipeline)."""
    sections = report.build_sections(entries, available)
    summary = report.summary_counts(
        sections,
        {
            "discovered": len(entries) + 1,
            "verified": len(entries),
            "unique": len(entries),
            "risk_excluded": 1,
        },
    )
    return write_excel_report(
        sections,
        summary,
        path=tmp_path / "superteam_report.xlsx",
        generated_at="2026-09-15 09:00 UTC",
        now=NOW,
    )


def available_from(entries: list[dict]) -> list[dict]:
    """Подходящие записи — тем же предикатом, что и в реальном прогоне."""
    return [entry for entry in entries if is_verified_open_listing(entry)]


def table_rows(book, sheet_name: str, columns: int) -> list:
    """Строки таблицы (данные начинаются с 5-й строки)."""
    sheet = book[sheet_name]
    return [
        row
        for row in sheet.iter_rows(min_row=5, max_col=columns)
        if row[0].value not in (None, "")
    ]


def link_targets(book, sheet_name: str, columns: int) -> list[str]:
    """Цели гиперссылок на листе (последняя колонка)."""
    sheet = book[sheet_name]
    targets = []
    for row in sheet.iter_rows(min_row=5, max_col=columns):
        if row[0].value in (None, ""):
            continue
        cell = sheet.cell(row=row[0].row, column=columns)
        targets.append(cell.hyperlink.target if cell.hyperlink else "")
    return targets


# --- AVAILABLE ------------------------------------------------------------------


def test_agent_allowed_verified_open_goes_to_available(tmp_path):
    entry = make_entry()
    book = load_workbook(build_workbook(tmp_path, [entry], available_from([entry])))
    rows = table_rows(book, SHEET_AVAILABLE, len(AVAILABLE_COLUMNS))

    assert len(rows) == 1
    row = rows[0]
    assert row[0].value == 1
    assert row[1].value == "HIGH"
    assert row[2].value == 82
    assert row[3].value == "Fix the CSV export bug"
    assert row[4].value == 500
    assert row[5].value == "USDC"
    assert row[8].value == "AGENT_ALLOWED"


def test_agent_only_verified_open_goes_to_available(tmp_path):
    entry = make_entry(agent_access="AGENT_ONLY", title="Agent-only task")
    book = load_workbook(build_workbook(tmp_path, [entry], available_from([entry])))
    rows = table_rows(book, SHEET_AVAILABLE, len(AVAILABLE_COLUMNS))
    assert [row[3].value for row in rows] == ["Agent-only task"]


# --- EXCLUDED / UNKNOWN ---------------------------------------------------------


def excluded_case(kind: str) -> dict:
    """Заведомо исключённые записи (как они приходят из pipeline)."""
    if kind == "human-only":
        return make_entry(
            slug="h",
            title="Human-only task",
            card_url="https://x/h",
            url="https://x/h",
            verification_status=VERIFIED_HUMAN_ONLY,
            agent_access="HUMAN_ONLY",
        )
    if kind == "winner":
        return make_entry(
            slug="w",
            title="Winner announced task",
            card_url="https://x/w",
            url="https://x/w",
            verification_status=VERIFIED_WINNER_ANNOUNCED,
            has_winners=True,
            final_decision="EXCLUDE",
            exclusion_reasons=["card verification: VERIFIED_WINNER_ANNOUNCED"],
        )
    if kind == "expired":
        return make_entry(
            slug="e",
            title="Expired task",
            card_url="https://x/e",
            url="https://x/e",
            verification_status=VERIFIED_EXPIRED,
            deadline=PAST,
            deadline_passed=True,
            final_decision="EXCLUDE",
            exclusion_reasons=["card verification: VERIFIED_EXPIRED", "card deadline already passed"],
        )
    if kind == "closed":
        return make_entry(
            slug="c",
            title="Closed task",
            card_url="https://x/c",
            url="https://x/c",
            verification_status=VERIFIED_CLOSED,
            final_decision="EXCLUDE",
            exclusion_reasons=["card verification: VERIFIED_CLOSED"],
        )
    if kind == "region":
        return make_entry(
            slug="r",
            title="Region restricted task",
            card_url="https://x/r",
            url="https://x/r",
            eligibility_status="REGION_RESTRICTED",
            final_decision="EXCLUDE",
            exclusion_reasons=["REGION_INELIGIBLE"],
        )
    if kind == "risk":
        return make_entry(
            slug="k",
            title="Risky task",
            card_url="https://x/k",
            url="https://x/k",
            financial_risk="HIGH",
            financial_risk_state="risk_detected",
            final_decision="EXCLUDE",
            exclusion_reasons=["REAL_FINANCIAL_ACTIVITY_REQUIRED"],
        )
    raise AssertionError(f"unknown case: {kind}")


@pytest.mark.parametrize("kind", ["human-only", "winner", "expired", "closed", "region", "risk"])
def test_excluded_cases_never_land_on_available(tmp_path, kind):
    entry = excluded_case(kind)
    book = load_workbook(build_workbook(tmp_path, [entry], available_from([entry])))

    assert table_rows(book, SHEET_AVAILABLE, len(AVAILABLE_COLUMNS)) == []
    rows = table_rows(book, SHEET_EXCLUDED, len(EXCLUDED_COLUMNS))
    assert [row[1].value for row in rows] == [entry["title"]]
    # причина показана по-человечески, а не техническим кодом
    assert rows[0][2].value
    assert "card verification:" not in str(rows[0][2].value)


def test_unknown_goes_to_unknown_sheet_and_never_to_available(tmp_path):
    entry = make_entry(
        slug="u",
        title="Unverifiable task",
        card_url="https://x/u",
        url="https://x/u",
        verification_status=UNKNOWN,
        deadline_confirmed=False,
        final_decision="EXCLUDE",
        exclusion_reasons=["card verification: UNKNOWN"],
        error="timeout after 20s",
    )
    book = load_workbook(build_workbook(tmp_path, [entry], available_from([entry])))

    assert table_rows(book, SHEET_AVAILABLE, len(AVAILABLE_COLUMNS)) == []
    assert table_rows(book, SHEET_EXCLUDED, len(EXCLUDED_COLUMNS)) == []
    rows = table_rows(book, SHEET_UNKNOWN, len(UNKNOWN_COLUMNS))
    assert [row[1].value for row in rows] == ["Unverifiable task"]
    assert "timeout" in str(rows[0][4].value)


# --- структура книги и ссылки ---------------------------------------------------


def test_workbook_has_expected_sheets_and_active_summary(tmp_path):
    entry = make_entry()
    book = load_workbook(build_workbook(tmp_path, [entry], available_from([entry])))

    assert book.sheetnames[:4] == [SHEET_SUMMARY, SHEET_AVAILABLE, SHEET_EXCLUDED, SHEET_UNKNOWN]
    assert book.active.title == SHEET_SUMMARY
    assert book[SHEET_SUMMARY]["A1"].value == "SUPERTEAM AGENT"
    assert book[SHEET_AVAILABLE].freeze_panes == "A5"
    assert book[SHEET_EXCLUDED].freeze_panes == "A5"


def test_available_hyperlinks_target_card_url(tmp_path):
    entry = make_entry()
    book = load_workbook(build_workbook(tmp_path, [entry], available_from([entry])))

    assert link_targets(book, SHEET_AVAILABLE, len(AVAILABLE_COLUMNS)) == [CARD_URL]
    cell = book[SHEET_AVAILABLE].cell(row=5, column=len(AVAILABLE_COLUMNS))
    assert cell.value == "Open card"
    assert CARD_URL not in str(cell.value)  # в ячейке короткий текст, не URL


def test_no_excluded_link_appears_on_available_sheet(tmp_path):
    excluded = excluded_case("winner")
    available_entry = make_entry()
    entries = [excluded, available_entry]
    book = load_workbook(build_workbook(tmp_path, entries, available_from(entries)))

    available_links = set(link_targets(book, SHEET_AVAILABLE, len(AVAILABLE_COLUMNS)))
    excluded_links = set(link_targets(book, SHEET_EXCLUDED, len(EXCLUDED_COLUMNS)))

    assert available_links == {CARD_URL}
    assert excluded_links == {"https://x/w"}
    assert not (available_links & excluded_links)


def test_excluded_sheet_sorted_agent_first_then_human_then_others():
    agent = report.build_item(make_entry(slug="a", agent_access="AGENT_ALLOWED", score=10, title="A"))
    human = report.build_item(make_entry(slug="b", agent_access="HUMAN_ONLY", score=90, title="B"))
    other = report.build_item(make_entry(slug="c", agent_access="UNKNOWN", score=99, title="C"))
    agent_low = report.build_item(make_entry(slug="d", agent_access="AGENT_ONLY", score=5, title="D"))

    ordered = sort_excluded([other, human, agent, agent_low])
    # агентские (score DESC) → human-only → остальные
    assert [item.title for item in ordered] == ["A", "D", "B", "C"]


def test_summary_counts_match_pipeline(tmp_path):
    available_entry = make_entry()
    winner = excluded_case("winner")
    unknown = make_entry(
        slug="u",
        title="U",
        card_url="https://x/u",
        url="https://x/u",
        verification_status=UNKNOWN,
        deadline_confirmed=False,
        final_decision="EXCLUDE",
        exclusion_reasons=["card verification: UNKNOWN"],
    )
    entries = [available_entry, winner, unknown]
    path = build_workbook(tmp_path, entries, available_from(entries))
    sheet = load_workbook(path)[SHEET_SUMMARY]
    values = {
        sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
        for row in range(1, 26)
    }

    assert values["Available for Agent"] == 1
    assert values["Excluded"] == 1
    assert values["Unknown"] == 1
    assert values["Verified"] == len(entries)
    assert values["Discovered"] == len(entries) + 1
    assert values["Risk Excluded"] == 1


def test_reward_split_and_real_deadline_datetime(tmp_path):
    entry = make_entry(reward="1 000 USDG", reward_amount=1000, reward_currency="USDG", token="USDG")
    book = load_workbook(build_workbook(tmp_path, [entry], available_from([entry])))
    row = table_rows(book, SHEET_AVAILABLE, len(AVAILABLE_COLUMNS))[0]

    assert row[4].value == 1000          # Reward Amount — число
    assert row[5].value == "USDG"        # Currency — отдельно
    assert row[6].value == datetime(2030, 1, 1, 0, 0)  # реальная дата (timezone-aware внутри)
    assert row[7].value == days_left(report.build_item(entry), NOW)


def test_urgent_deadline_is_highlighted(tmp_path):
    soon = (NOW + timedelta(days=1)).strftime("%Y-%m-%dT%H:%M:%SZ")
    entry = make_entry(deadline=soon)
    book = load_workbook(build_workbook(tmp_path, [entry], available_from([entry])))
    cell = book[SHEET_AVAILABLE].cell(row=5, column=8)  # Days Left

    assert isinstance(cell.value, (int, float))
    assert cell.value <= 3
    assert str(cell.fill.fgColor.rgb).endswith("FFC7CE")