"""Excel-отчёт: основной пользовательский результат (openpyxl).

Своей фильтрации здесь нет: лист ``AVAILABLE`` заполняется ТОЛЬКО из
``sections.available``, который сформирован существующим pipeline
(``is_verified_open_listing`` → card verification + eligibility + risk + scoring).
Исключённые и непроверяемые задания физически лежат на других листах, поэтому
ссылка на исключённое задание не может попасть на главный лист.

Листы: ``SUMMARY`` (dashboard), ``AVAILABLE``, ``EXCLUDED``, ``UNKNOWN``.
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Final, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from .config import (
    EXCEL_DESCRIPTION_LIMIT,
    EXCEL_REPORT_FILE,
    EXCEL_SOON_DAYS,
    EXCEL_URGENT_DAYS,
    RESULTS_FILE,
)
from .report import (
    ReportItem,
    ReportSections,
    exclusion_reason_short,
    generated_timestamp,
)
from .secrets import safe

#: Имена листов.
SHEET_SUMMARY: Final[str] = "SUMMARY"
SHEET_AVAILABLE: Final[str] = "AVAILABLE"
SHEET_EXCLUDED: Final[str] = "EXCLUDED"
SHEET_UNKNOWN: Final[str] = "UNKNOWN"

#: Колонки таблиц: (заголовок, ширина).
AVAILABLE_COLUMNS: Final[tuple[tuple[str, int], ...]] = (
    ("#", 5),
    ("Priority", 10),
    ("Score", 8),
    ("Task", 44),
    ("Reward Amount", 14),
    ("Currency", 9),
    ("Deadline (UTC)", 18),
    ("Days Left", 10),
    ("Agent Access", 15),
    ("Difficulty", 11),
    ("Risk", 11),
    ("Eligibility", 12),
    ("Region", 16),
    ("Submissions", 11),
    ("Type", 10),
    ("What to build", 50),
    ("Open Card", 12),
)
EXCLUDED_COLUMNS: Final[tuple[tuple[str, int], ...]] = (
    ("#", 5),
    ("Task", 42),
    ("Why excluded", 26),
    ("Reward Amount", 14),
    ("Currency", 9),
    ("Deadline (UTC)", 18),
    ("Verification", 24),
    ("Agent Access", 15),
    ("Risk", 11),
    ("Eligibility", 12),
    ("Score", 8),
    ("Final Decision", 14),
    ("Region", 16),
    ("Card", 12),
)
UNKNOWN_COLUMNS: Final[tuple[tuple[str, int], ...]] = (
    ("#", 5),
    ("Task", 42),
    ("Verification", 14),
    ("Reason", 26),
    ("Verification error", 40),
    ("Card", 12),
)

#: Стили (аккуратный «отчётный» вид, без внешних зависимостей).
TITLE_FONT: Final[Font] = Font(bold=True, size=16, color="1F3864")
SUB_FONT: Final[Font] = Font(bold=True, size=11, color="1F3864")
NOTE_FONT: Final[Font] = Font(size=10, color="595959")
HEADER_FONT: Final[Font] = Font(bold=True, color="FFFFFF")
HEADER_FILL: Final[PatternFill] = PatternFill("solid", fgColor="1F3864")
LINK_FONT: Final[Font] = Font(color="0563C1", underline="single")
_LINE: Final[Side] = Side(style="thin", color="BFBFBF")
BORDER: Final[Border] = Border(left=_LINE, right=_LINE, top=_LINE, bottom=_LINE)
WRAP_TOP: Final[Alignment] = Alignment(vertical="top", wrap_text=True)
CENTER: Final[Alignment] = Alignment(horizontal="center", vertical="center")
RIGHT: Final[Alignment] = Alignment(horizontal="right", vertical="top")
PRIORITY_FILLS: Final[dict[str, PatternFill]] = {
    "HIGH": PatternFill("solid", fgColor="FCE4D6"),
    "MEDIUM": PatternFill("solid", fgColor="FFF2CC"),
    "LOW": PatternFill("solid", fgColor="E2EFDA"),
}
PRIORITY_FONTS: Final[dict[str, Font]] = {
    "HIGH": Font(bold=True, color="C00000"),
    "MEDIUM": Font(bold=True, color="9C5700"),
    "LOW": Font(bold=True, color="375623"),
}
URGENT_FILL: Final[PatternFill] = PatternFill("solid", fgColor="FFC7CE")
SOON_FILL: Final[PatternFill] = PatternFill("solid", fgColor="FFEB9C")
DATE_FORMAT: Final[str] = "dd mmm yyyy hh:mm"
REWARD_FORMAT: Final[str] = "#,##0.##"


def _excel_datetime(moment: datetime | None) -> datetime | None:
    """Дата в виде, понятном Excel (UTC без tzinfo — Excel не хранит таймзоны)."""
    if moment is None:
        return None
    if moment.tzinfo is None:
        return moment
    return moment.astimezone(timezone.utc).replace(tzinfo=None)


def days_left(item: ReportItem, now: datetime) -> float | None:
    """Сколько дней осталось до подтверждённого дедлайна (``None`` — нет данных)."""
    moment = _excel_datetime(item.deadline_moment)
    if moment is None:
        return None
    return round((moment - _excel_datetime(now)).total_seconds() / 86400, 1)


def _reward_amount(item: ReportItem) -> float | None:
    """Числовая сумма награды (отдельная колонка, п.12 задания)."""
    raw = item.raw.get("reward_amount")
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return float(raw)
    return None


def _currency(item: ReportItem) -> str:
    return str(item.raw.get("reward_currency") or item.currency or "").upper()


def _topics(item: ReportItem) -> str:
    """Короткое описание «что нужно сделать» (обрезано, п.14 задания)."""
    text = " ".join(str(item.description_short or "").split())
    if not text:
        return ""
    if len(text) <= EXCEL_DESCRIPTION_LIMIT:
        return text
    return f"{text[:EXCEL_DESCRIPTION_LIMIT].rsplit(' ', 1)[0]}…"


def _write_table_header(
    sheet: Worksheet, columns: tuple[tuple[str, int], ...], *, title: str, subtitle: str
) -> int:
    """Записать заголовок листа и шапку таблицы; вернуть номер строки шапки."""
    sheet["A1"] = title
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = subtitle
    sheet["A2"].font = NOTE_FONT

    header_row = 4
    for index, (name, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        sheet.column_dimensions[cell.column_letter].width = width
    return header_row


def _style_body_cell(cell, *, alignment: Alignment = WRAP_TOP) -> None:
    """Единый вид ячеек тела таблицы (границы + выравнивание)."""
    cell.border = BORDER
    cell.alignment = alignment


def _set_link(cell, url: str, label: str = "Open card") -> None:
    """Настоящая Excel-гиперссылка (в ячейке — короткий текст, не URL)."""
    if not url:
        cell.value = "—"
        _style_body_cell(cell, alignment=CENTER)
        return
    cell.value = label
    cell.hyperlink = url
    cell.font = LINK_FONT
    _style_body_cell(cell, alignment=CENTER)


def _write_available_row(
    sheet: Worksheet, row: int, position: int, item: ReportItem, now: datetime
) -> None:
    """Одна строка главного листа (только задачи, прошедшие все фильтры)."""
    left = days_left(item, now)
    values: dict[str, Any] = {
        "#": position,
        "Priority": item.priority,
        "Score": item.score,
        "Task": safe(item.title),
        "Reward Amount": _reward_amount(item),
        "Currency": _currency(item),
        "Deadline (UTC)": _excel_datetime(item.deadline_moment),
        "Days Left": left,
        "Agent Access": item.agent_access,
        "Difficulty": item.difficulty or "UNKNOWN",
        "Risk": item.risk_status or "",
        "Eligibility": item.eligibility or "",
        "Region": safe(item.region),
        "Submissions": item.submissions,
        "Type": item.listing_type,
        "What to build": safe(_topics(item)),
    }
    for index, (name, _width) in enumerate(AVAILABLE_COLUMNS, start=1):
        cell = sheet.cell(row=row, column=index)
        if name == "Open Card":
            _set_link(cell, item.card_url)
            continue
        cell.value = values.get(name)
        _style_body_cell(cell, alignment=CENTER if name in ("#", "Submissions") else WRAP_TOP)
        if name == "Priority":
            cell.fill = PRIORITY_FILLS.get(item.priority.upper(), PatternFill())
            cell.font = PRIORITY_FONTS.get(item.priority.upper(), Font())
            cell.alignment = CENTER
        elif name == "Score":
            cell.alignment = RIGHT
            cell.number_format = "0"
        elif name == "Reward Amount":
            cell.alignment = RIGHT
            cell.number_format = REWARD_FORMAT
        elif name == "Deadline (UTC)":
            cell.alignment = CENTER
            if isinstance(cell.value, datetime):
                cell.number_format = DATE_FORMAT
        elif name == "Days Left":
            cell.alignment = CENTER
            cell.number_format = "0.0"
            if isinstance(left, (int, float)):
                if left <= EXCEL_URGENT_DAYS:
                    cell.fill = URGENT_FILL
                elif left <= EXCEL_SOON_DAYS:
                    cell.fill = SOON_FILL


def excluded_sort_key(item: ReportItem) -> tuple[int, float, str]:
    """Порядок листа EXCLUDED (п.15): агентские → human-only → остальные, score ↓.

    Так сразу видно, какие потенциально интересные задания были отброшены.
    """
    access = item.agent_access.upper()
    if access in ("AGENT_ONLY", "AGENT_ALLOWED"):
        group = 0
    elif access == "HUMAN_ONLY":
        group = 1
    else:
        group = 2
    try:
        score = -float(item.score or 0)
    except (TypeError, ValueError):
        score = 0.0
    return group, score, item.title.lower()


def sort_excluded(items: list[ReportItem]) -> list[ReportItem]:
    """Отсортировать исключённые записи для листа EXCLUDED."""
    return sorted(items, key=excluded_sort_key)


def _write_excluded_row(sheet: Worksheet, row: int, position: int, item: ReportItem) -> None:
    """Одна строка листа EXCLUDED (человеческая причина + ссылка для анализа)."""
    values: dict[str, Any] = {
        "#": position,
        "Task": safe(item.title),
        "Why excluded": exclusion_reason_short(item.exclusion_reasons) or "Filtered by pipeline",
        "Reward Amount": _reward_amount(item),
        "Currency": _currency(item),
        "Deadline (UTC)": _excel_datetime(item.deadline_moment),
        "Verification": item.verification_status,
        "Agent Access": item.agent_access,
        "Risk": item.risk_status or "",
        "Eligibility": item.eligibility or "",
        "Score": item.score,
        "Final Decision": item.decision,
        "Region": safe(item.region),
    }
    for index, (name, _width) in enumerate(EXCLUDED_COLUMNS, start=1):
        cell = sheet.cell(row=row, column=index)
        if name == "Card":
            _set_link(cell, item.card_url)
            continue
        cell.value = values.get(name)
        _style_body_cell(cell, alignment=CENTER if name in ("#", "Score", "Final Decision") else WRAP_TOP)
        if name == "Score":
            cell.alignment = RIGHT
            cell.number_format = "0"
        elif name == "Reward Amount":
            cell.alignment = RIGHT
            cell.number_format = REWARD_FORMAT
        elif name == "Deadline (UTC)":
            cell.alignment = CENTER
            if isinstance(cell.value, datetime):
                cell.number_format = DATE_FORMAT


def _write_unknown_row(sheet: Worksheet, row: int, position: int, item: ReportItem) -> None:
    """Одна строка листа UNKNOWN (непроверяемые задания)."""
    values: dict[str, Any] = {
        "#": position,
        "Task": safe(item.title),
        "Verification": item.verification_status,
        "Reason": exclusion_reason_short(item.exclusion_reasons) or "Card could not be verified",
        "Verification error": safe(item.error or "; ".join(item.evidence[:2])),
    }
    for index, (name, _width) in enumerate(UNKNOWN_COLUMNS, start=1):
        cell = sheet.cell(row=row, column=index)
        if name == "Card":
            _set_link(cell, item.card_url)
            continue
        cell.value = values.get(name)
        _style_body_cell(cell, alignment=CENTER if name == "#" else WRAP_TOP)


def _kv(sheet: Worksheet, row: int, label: str, value: Any, *, bold: bool = False) -> None:
    """Строка «метка → значение» на dashboard-листе."""
    label_cell = sheet.cell(row=row, column=1, value=label)
    label_cell.font = Font(bold=bold, size=11, color="1F3864" if bold else "000000")
    value_cell = sheet.cell(row=row, column=2, value=value)
    value_cell.font = Font(bold=bold, size=11, color="1F3864" if bold else "000000")
    value_cell.alignment = RIGHT


def _write_summary_sheet(
    sheet: Worksheet, summary: Mapping[str, int], *, generated_at: str, excel_name: str
) -> None:
    """Dashboard-лист: главные числа и разбивка (п.8 задания)."""
    sheet["A1"] = "SUPERTEAM AGENT"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = "Bounty Scan Report"
    sheet["A2"].font = SUB_FONT
    sheet["A4"] = "Generated:"
    sheet["A4"].font = NOTE_FONT
    sheet["B4"] = generated_at
    sheet["B4"].font = NOTE_FONT

    row = 6
    sheet.cell(row=row, column=1, value="RESULTS").font = SUB_FONT
    row += 1
    for label, key in (
        ("Available for Agent", "available"),
        ("Excluded", "excluded"),
        ("Unknown", "unknown"),
    ):
        _kv(sheet, row, label, summary.get(key, 0), bold=key == "available")
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="VERIFICATION").font = SUB_FONT
    row += 1
    for label, key in (("Verified", "verified"), ("Discovered", "discovered")):
        _kv(sheet, row, label, summary.get(key, 0))
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="BREAKDOWN").font = SUB_FONT
    row += 1
    for label, key in (
        ("Agent Allowed", "agent_allowed"),
        ("Human Only", "human_only"),
        ("Winner Announced", "winner_announced"),
        ("Risk Excluded", "risk_excluded"),
        ("Expired", "expired"),
    ):
        _kv(sheet, row, label, summary.get(key, 0))
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="FILES").font = SUB_FONT
    row += 1
    _kv(sheet, row, "Excel report", excel_name)
    _kv(sheet, row + 1, "JSON results", RESULTS_FILE.name)

    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 22
    sheet.sheet_view.showGridLines = False


def _finish_table(sheet: Worksheet, header_row: int, columns: int, data_rows: int) -> None:
    """Freeze panes, autofilter и печать «по-отчётному» (общая часть листов)."""
    from openpyxl.utils import get_column_letter

    last_column = get_column_letter(columns)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1).coordinate
    if data_rows > 0:
        sheet.auto_filter.ref = f"A{header_row}:{last_column}{header_row + data_rows}"
    sheet.page_setup.orientation = "landscape"
    sheet.print_title_rows = f"{header_row}:{header_row}"


def write_excel_report(
    sections: ReportSections,
    summary: Mapping[str, int],
    *,
    path: Path | str | None = None,
    generated_at: str = "",
    now: datetime | None = None,
) -> Path:
    """Собрать Excel-отчёт и сохранить его на диск.

    :param sections: разделы, построенные существующим pipeline
        (:func:`superteam_agent.report.build_sections`).
    :param summary: счётчики (:func:`superteam_agent.report.summary_counts`).
    :param path: куда сохранить (по умолчанию ``superteam_report.xlsx``).
    :param generated_at: метка времени отчёта.
    :param now: текущий момент (для «Days Left»; в тестах можно зафиксировать).
    :returns: путь к записанному файлу.
    """
    target = Path(path) if path is not None else EXCEL_REPORT_FILE
    moment = now or datetime.now(timezone.utc)
    stamp = generated_at or generated_timestamp()

    book = Workbook()
    summary_sheet = book.active
    summary_sheet.title = SHEET_SUMMARY
    _write_summary_sheet(summary_sheet, summary, generated_at=stamp, excel_name=target.name)

    available_sheet = book.create_sheet(SHEET_AVAILABLE)
    available_sheet.sheet_properties.tabColor = "375623"
    available_sheet.sheet_view.showGridLines = False
    header_row = _write_table_header(
        available_sheet,
        AVAILABLE_COLUMNS,
        title="AVAILABLE FOR AGENT",
        subtitle=(
            f"{len(sections.available)} заданий прошли все фильтры (карточка проверена, "
            f"deadline подтверждён, agent access разрешён, риск допустим). Generated: {stamp}"
        ),
    )
    for position, item in enumerate(sections.available, start=1):
        _write_available_row(available_sheet, header_row + position, position, item, moment)
    _finish_table(available_sheet, header_row, len(AVAILABLE_COLUMNS), len(sections.available))

    excluded_sheet = book.create_sheet(SHEET_EXCLUDED)
    excluded_sheet.sheet_properties.tabColor = "C00000"
    excluded_sheet.sheet_view.showGridLines = False
    header_row = _write_table_header(
        excluded_sheet,
        EXCLUDED_COLUMNS,
        title="EXCLUDED",
        subtitle=(
            f"{len(sections.excluded)} заданий исключены финальным pipeline "
            "(для ручного анализа; ссылки вторичны)"
        ),
    )
    excluded_items = sort_excluded(sections.excluded)
    for position, item in enumerate(excluded_items, start=1):
        _write_excluded_row(excluded_sheet, header_row + position, position, item)
    _finish_table(excluded_sheet, header_row, len(EXCLUDED_COLUMNS), len(excluded_items))

    unknown_sheet = book.create_sheet(SHEET_UNKNOWN)
    unknown_sheet.sheet_properties.tabColor = "BF8F00"
    unknown_sheet.sheet_view.showGridLines = False
    header_row = _write_table_header(
        unknown_sheet,
        UNKNOWN_COLUMNS,
        title="UNKNOWN / VERIFICATION FAILED",
        subtitle=(
            f"{len(sections.unknown)} заданий, у которых карточку проверить не удалось "
            "(такие задания НЕ считаются доступными)"
        ),
    )
    for position, item in enumerate(sections.unknown, start=1):
        _write_unknown_row(unknown_sheet, header_row + position, position, item)
    _finish_table(unknown_sheet, header_row, len(UNKNOWN_COLUMNS), len(sections.unknown))

    book.active = 0
    book.save(target)
    return target